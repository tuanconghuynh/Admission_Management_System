import io
import re
from datetime import date
from openpyxl import load_workbook
from pypdf import PdfReader
from app.db.session import SessionLocal
from app.models.applicant import Applicant, ApplicantDoc
from app.models.checklist import ChecklistItem, ChecklistVersion
from app.services.export_service import build_excel_bytes_by_items
from app.services.pdf_service import render_batch_pdf
from conftest import login


def seed():
    with SessionLocal.begin() as db:
        applicant = db.get(Applicant, '1234567890')
        applicant.ngay_nhan_hs = date(2026, 10, 2)
        version = db.query(ChecklistVersion).one()
        for index, (code, name, quantity) in enumerate([
            ('bang_tot_nghiep_dai_hoc', 'Bằng tốt nghiệp Đại học', 3),
            ('bang_diem_cao_dang', 'Bảng điểm Cao đẳng', 5),
            ('don_mien_giam', 'Đơn xin miễn giảm học phần', 2),
        ]):
            db.add(ChecklistItem(version_id=version.id, code=code, display_name=name, order_no=index))
            db.add(ApplicantDoc(applicant_ma_so_hv=applicant.ma_so_hv, code=code, display_name=name, so_luong=quantity))


def pdf_text(data):
    return '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)


def assert_quantities(text):
    for name, quantity in [('Bằng tốt nghiệp Đại học', 3), ('Bảng điểm Cao đẳng', 5), ('Đơn xin miễn giảm học phần', 2)]:
        assert re.search(rf'{quantity}\s+{name}', text) or re.search(rf'{name}\s+{quantity}', text), text
    assert 'HỒ SƠ XÉT MIỄN MÔN' not in text


def test_receipt_retains_quantities_and_has_no_exemption_block(client):
    seed()
    login(client)
    response = client.get('/applicants/1234567890/print')
    assert response.status_code == 200
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 1
    text = pdf_text(response.content)
    assert 'BIÊN NHẬN HỒ SƠ' in text
    assert_quantities(text)
    assert text.count('Bằng tốt nghiệp Đại học') == 2
    assert text.count('BIÊN NHẬN HỒ SƠ NHẬP HỌC') == 2
    assert text.count('Người nộp') == 2
    assert text.count('Người nhận') == 2
    assert 'Mã hồ sơ: 0001' in text


def test_folder_cover_and_legacy_alias(client):
    seed()
    login(client)
    for path in ('/applicants/1234567890/folder-cover', '/applicants/1234567890/postal-print'):
        response = client.get(path)
        assert response.status_code == 200, response.text
        text = pdf_text(response.content)
        assert 'BIÊN NHẬN HỒ SƠ NHẬP HỌC' in text and 'MÃ HS : 0001' in text
        assert 'BÌA HỒ SƠ' not in text
        assert 'Danh mục hồ sơ' in text
        assert_quantities(text)


def test_excel_single_sheet_keeps_received_quantities(client):
    seed()
    login(client)
    response = client.get('/api/export/excel?date=02/10/2026')
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ['Hồ sơ']
    rows = list(book.active.values)
    for name, quantity in [('Bằng tốt nghiệp Đại học', 3), ('Bảng điểm Cao đẳng', 5), ('Đơn xin miễn giảm học phần', 2)]:
        assert rows[1][rows[0].index(name)] == quantity


def test_batch_cover_uses_each_applicants_documents(client):
    seed()
    with SessionLocal.begin() as db:
        version = db.query(ChecklistVersion).one()
        db.add(Applicant(ma_so_hv='1234567891', ma_ho_so='0002', ho_ten='Trần Thị Bình', checklist_version_id=version.id))
        db.flush()
        db.add(ApplicantDoc(applicant_ma_so_hv='1234567891', code='bang_tot_nghiep_dai_hoc', so_luong=7))
    with SessionLocal() as db:
        applicants = db.query(Applicant).order_by(Applicant.ma_so_hv).all()
        items = db.query(ChecklistItem).all()
        documents = db.query(ApplicantDoc).all()
        output = render_batch_pdf(applicants, items, documents, print_type='COVER')
        text = pdf_text(output)
        assert re.search(r'Bằng tốt nghiệp Đại học\s+3', text)
        assert re.search(r'Bằng tốt nghiệp Đại học\s+7', text)
